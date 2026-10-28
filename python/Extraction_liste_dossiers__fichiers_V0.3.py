"""Extraction de l'arborescence d'un dossier vers Excel - V0.3 (2026-10-28)

Même usage et même mise en page que la V0.2 :
*   Fenêtre de sélection du dossier racine, puis fenêtre d'options.
*   Une colonne par niveau, liens cliquables vers les sous-dossiers, taille et date de modification.
*   Regroupement des fichiers nombreux d'un même type ("[......] 12 fichiers .pdf").

Nouveautés V0.3 (robustesse) :
*   Un dossier inaccessible (droits, réseau) n'arrête plus l'extraction : il est signalé dans la colonne Remarque.
*   Les liens symboliques et jonctions Windows ne sont plus suivis (plus de boucles ni de doublons).
*   Parcours avec os.scandir : 1 appel système par élément au lieu de 4 (beaucoup plus rapide sur réseau).
*   Tri alphabétique naturel (A.txt avant b.txt, plan_2 avant plan_10).
*   Taille en octets et date réelles : triables et filtrables dans Excel. Choix de l'emplacement du fichier Excel.
*   Un nom de fichier commençant par "=" n'est plus interprété comme une formule.
*   Dossiers au-delà de la profondeur max signalés ("Contenu non exploré"), chemins > 260 caractères signalés.
*   Feuille "Résumé" : nombre de dossiers et fichiers, taille totale, éléments regroupés, ignorés ou inaccessibles.
*   Plus de dépendance à pandas : seul openpyxl est requis.
"""
import os
import re
import sys
import time
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

VERSION = "V0.3 (2026-10-28)"
AUTEUR = "Anthony Chéruel — acheruel.github.io/Arborescento"

# --- OPTIONS PAR DÉFAUT ---
PROFONDEUR_MAX_DEFAUT = 7
PROFONDEUR_MAX_LIMITE = 50
SEUIL_PDF_DEFAUT = 10

# Fichiers ignorés complètement (insensible à la casse) : noms exacts, préfixes et extensions
FICHIERS_A_IGNORER = {"thumbs.db", ".ds_store", "desktop.ini"}
PREFIXES_A_IGNORER = ("~$", ".~lock.")  # temporaires Office et LibreOffice
EXTENSIONS_A_IGNORER = {".log", ".tmp"}

# Regroupement : à partir de combien de fichiers d'une même extension dans un dossier on les regroupe sur une ligne
AUTRES_SEUILS = {
    ".dwg": 5, ".dcn": 2, ".jpg": 5, ".jpeg": 5, ".mov": 1, ".png": 5, ".docx": 8, ".doc": 8,
    ".xls": 7, ".xlsx": 7, ".xlsm": 7, ".txt": 2, ".dgn": 2, ".tif": 2, ".tiff": 2, ".zip": 5,
    ".heic": 2, ".ctb": 2,
}

# Mise en forme "chaud / froid" de la date de modification
JOURS_RECENT = 15   # fond vert si modifié il y a moins de 15 jours
JOURS_ANCIEN = 365  # gris italique si modifié il y a plus d'un an

LIMITE_CHEMIN_WINDOWS = 260
LIMITE_LIGNES_EXCEL = 1_048_576
FILE_ATTRIBUTE_REPARSE_POINT = 0x400


def handle_long_path(path):
    r"""
    Contournement de la limite MAX_PATH de 260 caractères sous Windows.
    Gère les chemins locaux (C:\...) et réseau UNC (\\serveur\...).
    """
    if os.name != 'nt':
        return path
    path = os.path.abspath(path)
    if path.startswith('\\\\?\\'):
        return path
    if path.startswith('\\\\'):
        return f"\\\\?\\UNC\\{path[2:]}"
    return f"\\\\?\\{path}"


def formater_taille(taille_en_octets):
    """Convertit une taille en octets en texte lisible (o, Ko, Mo, Go, To)."""
    if taille_en_octets is None:
        return ""
    unites = ['o', 'Ko', 'Mo', 'Go', 'To']
    taille = float(taille_en_octets)
    n = 0
    while taille >= 1024 and n < len(unites) - 1:
        taille /= 1024
        n += 1
    # Pas de décimale pour 'o' et 'Ko', une seule décimale pour le reste
    return f"{taille:.0f} {unites[n]}" if n <= 1 else f"{taille:.1f} {unites[n]}"


def cle_tri_naturel(nom):
    """Tri alphabétique insensible à la casse, avec les nombres dans l'ordre (plan_2 avant plan_10)."""
    return [int(t) if t.isdigit() else t for t in re.split(r'(\d+)', nom.casefold())]


def est_ignore(nom):
    nom_min = nom.lower()
    return (nom_min in FICHIERS_A_IGNORER
            or nom_min.startswith(PREFIXES_A_IGNORER)
            or os.path.splitext(nom_min)[1] in EXTENSIONS_A_IGNORER)


def est_lien(entree):
    """Vrai pour un lien symbolique ou une jonction Windows : on ne les suit pas pour éviter les boucles."""
    try:
        if entree.is_symlink():
            return True
        if os.name == 'nt':
            attributs = entree.stat(follow_symlinks=False).st_file_attributes
            return bool(attributs & FILE_ATTRIBUTE_REPARSE_POINT) and entree.is_dir(follow_symlinks=False)
    except OSError:
        pass
    return False


def date_modif(stat_result):
    try:
        return datetime.fromtimestamp(stat_result.st_mtime)
    except (OSError, OverflowError, ValueError):
        return None


def nouvelle_ligne(type_ligne, profondeur, nom, chemin=None, taille=None, date=None, remarque=""):
    return {"type": type_ligne, "profondeur": profondeur, "nom": nom, "chemin": chemin,
            "taille": taille, "date": date, "remarque": remarque}


def lister_dossier(chemin, stats):
    """Retourne (fichiers, dossiers) triés sous forme de listes d'os.DirEntry, ou None si inaccessible."""
    try:
        with os.scandir(handle_long_path(chemin)) as it:
            entrees = list(it)
    except OSError as e:
        stats["inaccessibles"].append((chemin, e.strerror or type(e).__name__))
        return None
    fichiers, dossiers = [], []
    for e in entrees:
        try:
            if e.is_dir(follow_symlinks=False) or (e.is_symlink() and e.is_dir()):
                dossiers.append(e)
            elif e.is_file(follow_symlinks=False) or e.is_symlink():
                fichiers.append(e)
        except OSError:
            fichiers.append(e)
    fichiers.sort(key=lambda e: cle_tri_naturel(e.name))
    dossiers.sort(key=lambda e: cle_tri_naturel(e.name))
    return fichiers, dossiers


def parcourir_en_profondeur(chemin_actuel, profondeur, lignes, options, stats):
    """Parcourt récursivement chemin_actuel et ajoute une ligne par fichier, groupe ou sous-dossier."""
    contenu = lister_dossier(chemin_actuel, stats)
    if contenu is None:
        return False
    fichiers, dossiers = contenu

    candidats = []
    for f in fichiers:
        if est_ignore(f.name):
            stats["ignores"] += 1
        else:
            candidats.append(f)

    exclus = set()
    if options['activer']:
        groupes = {}
        for f in candidats:
            groupes.setdefault(os.path.splitext(f.name)[1].lower(), []).append(f)
        seuils = dict(options['autres_seuils'])
        seuils['.pdf'] = options['seuil_pdf']
        for ext, membres in groupes.items():
            seuil = seuils.get(ext)
            if seuil is None or len(membres) < max(seuil, 1):
                continue
            tailles, dates = [], []
            for f in membres:
                try:
                    st = f.stat(follow_symlinks=False)
                    tailles.append(st.st_size)
                    d = date_modif(st)
                    if d:
                        dates.append(d)
                except OSError:
                    pass
            nom = (f"[......] {len(membres)} fichiers {ext or '(sans extension)'} : "
                   f"{membres[0].name} … {membres[-1].name} (regroupés car ≥ {seuil})")
            lignes.append(nouvelle_ligne("groupe", profondeur, nom, taille=sum(tailles),
                                         date=max(dates) if dates else None))
            stats["fichiers"] += len(membres)
            stats["regroupes"] += len(membres)
            stats["octets"] += sum(tailles)
            exclus.update(f.name for f in membres)

    for f in candidats:
        if f.name in exclus:
            continue
        chemin = os.path.join(chemin_actuel, f.name)
        remarque = ""
        try:
            st = f.stat(follow_symlinks=False)
            taille, date = st.st_size, date_modif(st)
            stats["octets"] += taille
        except OSError:
            taille, date, remarque = None, None, "Inaccessible"
        if f.is_symlink():
            remarque = "Lien symbolique"
        if len(chemin) > LIMITE_CHEMIN_WINDOWS:
            remarque = (remarque + " · " if remarque else "") + f"Chemin > {LIMITE_CHEMIN_WINDOWS} car."
            stats["chemins_longs"] += 1
        lignes.append(nouvelle_ligne("fichier", profondeur, f.name, chemin, taille, date, remarque))
        stats["fichiers"] += 1

    for d in dossiers:
        chemin = os.path.join(chemin_actuel, d.name)
        try:
            date = date_modif(d.stat(follow_symlinks=False))
        except OSError:
            date = None
        ligne = nouvelle_ligne("dossier", profondeur, d.name, chemin, date=date)
        lignes.append(ligne)
        stats["dossiers"] += 1
        if len(chemin) > LIMITE_CHEMIN_WINDOWS:
            ligne["remarque"] = f"Chemin > {LIMITE_CHEMIN_WINDOWS} car."
            stats["chemins_longs"] += 1
        if est_lien(d):
            ligne["remarque"] = "Lien ou jonction : non suivi"
            stats["liens"] += 1
        elif profondeur + 1 >= options['profondeur_max']:
            # On ne descend pas plus loin, mais on signale s'il reste du contenu
            try:
                with os.scandir(handle_long_path(chemin)) as it:
                    if next(it, None) is not None:
                        ligne["remarque"] = "Contenu non exploré (profondeur max)"
                        stats["non_explores"] += 1
            except OSError as e:
                ligne["remarque"] = "Accès refusé"
                stats["inaccessibles"].append((chemin, e.strerror or type(e).__name__))
        elif not parcourir_en_profondeur(chemin, profondeur + 1, lignes, options, stats):
            ligne["remarque"] = "Accès refusé"
    return True


def ecrire_texte(cellule, texte):
    """Écrit du texte brut : un nom commençant par '=' ne doit jamais devenir une formule."""
    cellule.value = texte
    cellule.data_type = 's'


def generer_rapport_excel(chemin_racine, fichier_sortie, profondeur_max, appliquer_filtre, seuil_pdf,
                          autres_seuils=AUTRES_SEUILS, reessayer_si_verrouille=None):
    """
    Parcourt chemin_racine et écrit le fichier Excel. Retourne le dictionnaire de statistiques.
    reessayer_si_verrouille : fonction appelée si le fichier de sortie est verrouillé (ouvert dans Excel) ;
    si elle retourne True, l'écriture est retentée sans refaire le parcours.
    """
    if not os.path.isdir(handle_long_path(chemin_racine)):
        raise NotADirectoryError(f"Le dossier '{chemin_racine}' n'est pas valide.")
    chemin_racine = os.path.normpath(os.path.abspath(chemin_racine))

    options = {'profondeur_max': profondeur_max, 'activer': appliquer_filtre,
               'seuil_pdf': seuil_pdf, 'autres_seuils': autres_seuils}
    stats = {"dossiers": 0, "fichiers": 0, "octets": 0, "regroupes": 0, "ignores": 0, "liens": 0,
             "non_explores": 0, "chemins_longs": 0, "inaccessibles": []}
    lignes = []
    if not parcourir_en_profondeur(chemin_racine, 0, lignes, options, stats):
        raise PermissionError(f"Accès refusé au dossier racine '{chemin_racine}'.")

    maintenant = datetime.now()
    descriptif = [
        "Ce fichier Excel contient une liste brute de l'arborescence des sous-dossiers et fichiers du dossier sélectionné suivant :",
        f"Dossier racine : {chemin_racine}",
        f"Date de génération arborescence : {maintenant.strftime('%Y-%m-%d %H:%M')}",
        f"Filtre appliqué sur le nombre de fichier par sous-dossier : {'Oui' if appliquer_filtre else 'Non'}",
        f"Profondeur d'analyse max : {profondeur_max}",
        "",
        f"Script Python {VERSION} par {AUTEUR}",
        "",
    ]
    ligne_entete = len(descriptif) + 1
    max_donnees = LIMITE_LIGNES_EXCEL - ligne_entete
    tronque = len(lignes) > max_donnees
    if tronque:
        lignes = lignes[:max_donnees]

    # Colonnes : Niveau 1..N | Taille | Dernière modif. | Octets | Remarque
    col_taille = profondeur_max + 1
    col_date, col_octets, col_remarque = col_taille + 1, col_taille + 2, col_taille + 3

    wb = Workbook()
    ws = wb.active
    ws.title = "Arborescence"
    ws.sheet_view.showGridLines = False
    for i, texte in enumerate(descriptif, start=1):
        ecrire_texte(ws.cell(row=i, column=1), texte)

    police_entete = Font(bold=True)
    entetes = [f'Niveau {i + 1}' for i in range(profondeur_max)] + ['Taille', 'Dernière modif.', 'Octets', 'Remarque']
    for c, texte in enumerate(entetes, start=1):
        ws.cell(row=ligne_entete, column=c, value=texte).font = police_entete

    police_dossier = Font(bold=True, underline="single", color="000080")
    police_groupe = Font(italic=True, color="595959")
    police_ancien = Font(color="808080", italic=True)
    police_remarque = Font(color="C00000")
    fond_recent = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")

    for r, ligne in enumerate(lignes, start=ligne_entete + 1):
        cellule_nom = ws.cell(row=r, column=ligne["profondeur"] + 1)
        ecrire_texte(cellule_nom, ligne["nom"])
        if ligne["type"] == "dossier":
            cellule_nom.hyperlink = ligne["chemin"]
            cellule_nom.font = police_dossier
            ws.cell(row=r, column=col_taille, value='---')
        else:
            if ligne["type"] == "groupe":
                cellule_nom.font = police_groupe
            if ligne["taille"] is not None:
                ws.cell(row=r, column=col_taille, value=formater_taille(ligne["taille"]))
                ws.cell(row=r, column=col_octets, value=ligne["taille"]).number_format = '#,##0'
        if ligne["date"] is not None:
            cellule_date = ws.cell(row=r, column=col_date, value=ligne["date"])
            cellule_date.number_format = 'yyyy-mm-dd'
            age = (maintenant - ligne["date"]).days
            if age < JOURS_RECENT:
                cellule_date.fill = fond_recent
            elif age > JOURS_ANCIEN:
                cellule_date.font = police_ancien
        if ligne["remarque"]:
            cellule = ws.cell(row=r, column=col_remarque)
            ecrire_texte(cellule, ligne["remarque"])
            cellule.font = police_remarque

    ws.column_dimensions[get_column_letter(profondeur_max)].width = 85
    ws.column_dimensions[get_column_letter(col_taille)].width = 11.5
    ws.column_dimensions[get_column_letter(col_date)].width = 17
    ws.column_dimensions[get_column_letter(col_octets)].width = 15
    ws.column_dimensions[get_column_letter(col_remarque)].width = 36
    ws.freeze_panes = f'A{ligne_entete + 1}'

    resume = wb.create_sheet("Résumé")
    resume.sheet_view.showGridLines = False
    valeurs = [
        ("Dossier racine", chemin_racine),
        ("Date de génération", maintenant.strftime('%Y-%m-%d %H:%M')),
        ("Dossiers", stats["dossiers"]),
        ("Fichiers (y compris regroupés)", stats["fichiers"]),
        ("Taille totale", formater_taille(stats["octets"])),
        ("Fichiers regroupés", stats["regroupes"]),
        ("Fichiers ignorés (temporaires, Thumbs.db, .log…)", stats["ignores"]),
        ("Liens et jonctions non suivis", stats["liens"]),
        ("Dossiers non explorés (profondeur max)", stats["non_explores"]),
        (f"Chemins > {LIMITE_CHEMIN_WINDOWS} caractères", stats["chemins_longs"]),
        ("Dossiers inaccessibles", len(stats["inaccessibles"])),
    ]
    if tronque:
        valeurs.append(("ATTENTION", f"Liste tronquée à {max_donnees} lignes (limite Excel)"))
    for i, (cle, val) in enumerate(valeurs, start=1):
        resume.cell(row=i, column=1, value=cle).font = police_entete
        if isinstance(val, str):
            ecrire_texte(resume.cell(row=i, column=2), val)
        else:
            resume.cell(row=i, column=2, value=val)
    if stats["inaccessibles"]:
        debut = len(valeurs) + 2
        resume.cell(row=debut, column=1, value="Dossiers inaccessibles").font = police_entete
        for j, (chemin, raison) in enumerate(stats["inaccessibles"], start=debut + 1):
            ecrire_texte(resume.cell(row=j, column=1), chemin)
            ecrire_texte(resume.cell(row=j, column=2), raison)
    resume.column_dimensions['A'].width = 48
    resume.column_dimensions['B'].width = 60

    while True:
        try:
            wb.save(fichier_sortie)
            break
        except PermissionError:
            if reessayer_si_verrouille is None or not reessayer_si_verrouille(fichier_sortie):
                raise
    stats["tronque"] = tronque
    return stats


# ---------------- Interface (tkinter) ----------------

def activer_dpi_windows():
    """Évite une interface floue sur les écrans haute résolution Windows."""
    if os.name == 'nt':
        try:
            import ctypes
            ctypes.windll.shcore.SetProcessDpiAwareness(1)
        except Exception:
            pass


def choisir_dossier_racine():
    import tkinter as tk
    from tkinter import filedialog
    root = tk.Tk()
    root.withdraw()
    print("Veuillez sélectionner le dossier racine à analyser...")
    chemin = filedialog.askdirectory(title="Sélectionnez le dossier racine")
    root.destroy()
    return chemin


def choisir_fichier_sortie(dossier_racine):
    import tkinter as tk
    from tkinter import filedialog
    root = tk.Tk()
    root.withdraw()
    nom_racine = os.path.basename(os.path.normpath(dossier_racine)) or "racine"
    nom_racine = re.sub(r'[<>:"/\\|?*]', '_', nom_racine)
    chemin = filedialog.asksaveasfilename(
        title="Enregistrer le fichier Excel",
        initialfile=f"Arborescence_{nom_racine}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
        defaultextension=".xlsx", filetypes=[("Classeur Excel", "*.xlsx")])
    root.destroy()
    return chemin


def demander_reessai(fichier_sortie):
    import tkinter as tk
    from tkinter import messagebox
    root = tk.Tk()
    root.withdraw()
    reponse = messagebox.askretrycancel("Fichier verrouillé", f"Impossible d'écrire '{fichier_sortie}'.\n"
                                        "Fermez-le dans Excel, puis cliquez sur Réessayer.")
    root.destroy()
    return reponse


def message(titre, texte, erreur=False):
    import tkinter as tk
    from tkinter import messagebox
    root = tk.Tk()
    root.withdraw()
    (messagebox.showerror if erreur else messagebox.showinfo)(titre, texte)
    root.destroy()


def configurer_options_utilisateur(defaut_filtre=True, defaut_seuil_pdf=SEUIL_PDF_DEFAUT,
                                   defaut_profondeur=PROFONDEUR_MAX_DEFAUT):
    """
    Fenêtre de configuration. Retourne (appliquer_filtre, seuil_pdf, profondeur_max), ou None si annulé.
    """
    import tkinter as tk
    from tkinter import messagebox
    root = tk.Tk()
    root.title("Configuration de l'Extraction")
    w, h = 420, 300
    x, y = (root.winfo_screenwidth() - w) / 2, (root.winfo_screenheight() - h) / 2
    root.geometry(f'{w}x{h}+{int(x)}+{int(y)}')
    root.configure(bg="#f0f0f0")
    font_std = ("Segoe UI", 10)
    font_bold = ("Segoe UI", 10, "bold")

    var_filtre = tk.BooleanVar(value=defaut_filtre)
    var_seuil = tk.StringVar(value=str(defaut_seuil_pdf))
    var_profondeur = tk.StringVar(value=str(defaut_profondeur))

    frame = tk.Frame(root, bg="#f0f0f0", padx=20, pady=20)
    frame.pack(fill="both", expand=True)
    tk.Label(frame, text="Veuillez configurer les paramètres de l'extraction :",
             bg="#f0f0f0", font=font_bold, anchor="w").pack(fill="x", pady=(0, 15))

    frame_prof = tk.Frame(frame, bg="#f0f0f0")
    frame_prof.pack(fill="x", pady=(0, 10))
    tk.Label(frame_prof, text="Profondeur max. d'analyse (niveaux) :", bg="#f0f0f0", font=font_std).pack(side="left")
    tk.Entry(frame_prof, textvariable=var_profondeur, width=8, font=font_std).pack(side="left", padx=(10, 0))

    frame_seuil = tk.Frame(frame, bg="#f0f0f0")
    lbl_seuil = tk.Label(frame_seuil, text="Regrouper les PDF à partir de (fichiers) :", bg="#f0f0f0", font=font_std)
    entry_seuil = tk.Entry(frame_seuil, textvariable=var_seuil, width=8, font=font_std)

    def toggle_state():
        etat = "normal" if var_filtre.get() else "disabled"
        entry_seuil.config(state=etat)
        lbl_seuil.config(state=etat, fg="black" if var_filtre.get() else "gray")

    tk.Checkbutton(frame, text="Appliquer le filtre de regroupement", variable=var_filtre, command=toggle_state,
                   bg="#f0f0f0", font=font_std, anchor="w").pack(fill="x", pady=(0, 10))
    frame_seuil.pack(fill="x", pady=(0, 20), padx=20)
    lbl_seuil.pack(side="left")
    entry_seuil.pack(side="left", padx=(10, 0))
    toggle_state()

    frame_btn = tk.Frame(frame, bg="#f0f0f0")
    frame_btn.pack(fill="x", side="bottom")
    resultats = {}

    def valider():
        try:
            seuil, prof = int(var_seuil.get()), int(var_profondeur.get())
            if seuil < 1 or not 1 <= prof <= PROFONDEUR_MAX_LIMITE:
                raise ValueError
        except ValueError:
            messagebox.showerror("Erreur de format",
                                 f"Seuil : entier ≥ 1. Profondeur : entier entre 1 et {PROFONDEUR_MAX_LIMITE}.")
            return
        resultats.update(filtre=var_filtre.get(), seuil=seuil, profondeur=prof)
        root.destroy()

    tk.Button(frame_btn, text="Lancer l'extraction", command=valider, bg="#0078d7", fg="white",
              font=font_bold, relief="flat", padx=10, pady=5).pack(side="right")
    tk.Button(frame_btn, text="Annuler", command=root.destroy, bg="#e0e0e0",
              font=font_std, relief="flat", padx=10, pady=5).pack(side="right", padx=(0, 10))
    root.protocol("WM_DELETE_WINDOW", root.destroy)
    root.bind("<Return>", lambda _e: valider())
    root.mainloop()

    if not resultats:
        return None
    return resultats["filtre"], resultats["seuil"], resultats["profondeur"]


def main():
    activer_dpi_windows()
    dossier_racine = choisir_dossier_racine()
    if not dossier_racine:
        print("\nAucun dossier sélectionné. Arrêt du programme.")
        return 1
    options = configurer_options_utilisateur()
    if options is None:
        print("Annulation par l'utilisateur.")
        return 1
    appliquer_filtre, seuil_pdf, profondeur_max = options

    fichier_sortie = choisir_fichier_sortie(dossier_racine)
    if not fichier_sortie:
        print("Aucun fichier de sortie choisi. Arrêt du programme.")
        return 1

    debut = time.time()
    try:
        stats = generer_rapport_excel(dossier_racine, fichier_sortie, profondeur_max, appliquer_filtre, seuil_pdf,
                                      reessayer_si_verrouille=demander_reessai)
    except Exception as e:
        message("Erreur", f"Une erreur est survenue : {e}", erreur=True)
        return 1
    duree = time.time() - debut

    bilan = (f"{stats['dossiers']} dossiers, {stats['fichiers']} fichiers, {formater_taille(stats['octets'])}"
             f" — {duree:.1f} s")
    if stats["inaccessibles"]:
        bilan += f"\n{len(stats['inaccessibles'])} dossier(s) inaccessible(s) : voir la feuille Résumé."
    if stats["tronque"]:
        bilan += "\nATTENTION : liste tronquée à la limite de lignes d'Excel."
    print(f"\nLe fichier Excel '{fichier_sortie}' a été généré.\n{bilan}")
    message("Extraction terminée", f"{fichier_sortie}\n\n{bilan}")
    if os.name == 'nt':
        try:
            os.startfile(fichier_sortie)
        except OSError:
            pass
    return 0


if __name__ == '__main__':
    sys.exit(main())
